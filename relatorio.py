# -*- coding: utf-8 -*-
"""
Gera o relatório dos testes do registro de presença para o TCC.

O conteúdo é montado uma única vez (lista de blocos: títulos, parágrafos,
tabelas e figuras) e depois renderizado em dois formatos:
    - .docx no padrão ABNT (gerar_docx)
    - seção LaTeX para o template de artigos da SBC, em .zip com as figuras
      (gerar_latex)

Chamado pelos botões da página de testes, mas também pode ser rodado direto
(usa relatorio_config.json e testes.json):
    python relatorio.py          # gera o .docx
    python relatorio.py latex    # gera o .zip com a seção LaTeX
"""

import io
import os
import re
import time
import zipfile

from cadastrar import BASE_DIR
from casos_teste import (
    CASOS, VISITANTE, NOMES_DECISAO, calcular_metricas, metricas_por,
    carregar_resultados,
)

DOCX_PATH = os.path.join(BASE_DIR, "relatorio_tcc.docx")
LATEX_PATH = os.path.join(BASE_DIR, "relatorio_tcc_latex.zip")

# referências no texto: [[tab:id]] e [[fig:id]] viram "Tabela 2" no .docx
# e "Tabela~\ref{tab:id}" no LaTeX
REF = re.compile(r"\[\[(tab|fig):([\w-]+)\]\]")


def pct(v):
    return "–" if v is None else f"{v:.1f}%".replace(".", ",")


def num(v, casas=2):
    return "–" if v is None else f"{v:.{casas}f}".replace(".", ",")


def _data_br(iso):
    a, m, d = iso.split("-")
    return f"{d}/{m}/{a}"


# ---------------------------------------------------------------- gráficos

def _grafico_barras(rotulos, series, titulo_y):
    """series: lista de (nome, valores, cor). Retorna o PNG em bytes."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(8, 3.6))
    largura = 0.8 / len(series)
    for i, (nome, valores, cor) in enumerate(series):
        xs = [x + (i - (len(series) - 1) / 2) * largura
              for x in range(len(rotulos))]
        barras = ax.bar(xs, [v or 0 for v in valores], largura,
                        label=nome, color=cor)
        for b, v in zip(barras, valores):
            if v is not None:
                ax.annotate(f"{v:.0f}",
                            (b.get_x() + b.get_width() / 2, b.get_height()),
                            ha="center", va="bottom", fontsize=7,
                            xytext=(0, 2), textcoords="offset points")
    ax.set_xticks(range(len(rotulos)))
    ax.set_xticklabels(rotulos)
    ax.set_ylabel(titulo_y)
    ax.set_ylim(0, 112)
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", alpha=.3)
    if len(series) > 1:
        ax.legend(frameon=False, fontsize=8)
    fig.tight_layout()

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


# ---------------------------------------------------------------- conteúdo

def montar_conteudo(config, parametros, resultados, pessoas):
    """Retorna a lista de blocos do relatório.

    Blocos:
        ("titulo", nivel, texto)
        ("par", texto)
        ("itens", [(rotulo, texto), ...])
        ("tabela", id, titulo, cabecalho, linhas, larguras_cm)
        ("figura", id, titulo, png_bytes)
        ("apendice",) - daqui em diante só entra no .docx
    """
    b = []
    titulo = lambda texto, nivel=1: b.append(("titulo", nivel, texto))
    par = lambda texto: b.append(("par", texto))

    # anonimização: P1..Pn na ordem alfabética dos cadastrados
    cadastrados = sorted(pessoas)
    if config.get("anonimizar"):
        apelidos = {n: f"P{i}" for i, n in enumerate(cadastrados, 1)}
    else:
        apelidos = {n: n for n in cadastrados}

    def nome(n):
        return "–" if n is None else apelidos.get(n, n)

    geral = calcular_metricas(resultados)
    por_caso = metricas_por(resultados, "caso")
    por_participante = metricas_por(resultados, "participante")
    casos_feitos = [c for c in CASOS if c["id"] in por_caso]
    genuinos = [c for c in casos_feitos if c["esperado"] == "liberado"]
    fraudes = [c for c in casos_feitos if c["esperado"] == "negado"]

    # ------------------------------------------------------------ objetivo
    titulo("Objetivo dos testes")
    par("Os testes descritos a seguir foram realizados com o protótipo de "
        "registro automático de presença por reconhecimento facial "
        "desenvolvido neste trabalho. O protótipo simula a entrada de uma "
        "sala de aula: uma câmera instalada na porta captura o rosto de quem "
        "entra e, ao reconhecer um aluno ou professor cadastrado, registra "
        "automaticamente sua presença na lista do dia; quando o rosto não é "
        "reconhecido, nenhuma presença é registrada. Os testes buscam medir a "
        "confiabilidade do registro em diferentes distâncias e ângulos do "
        "rosto em relação à câmera e verificar a resistência do sistema a "
        "fraudes, em especial à tentativa de um aluno registrar a presença "
        "de um colega ausente apresentando uma foto ou um vídeo dele à câmera.")

    # ------------------------------------------------------------ metodologia
    titulo("Metodologia")
    titulo("Protótipo", 2)
    par("O protótipo foi implementado em Python com a biblioteca OpenCV e "
        "possui uma interface web desenvolvida com o framework Flask, com "
        "telas de cadastro, de registro de presença e de testes. A detecção "
        "do rosto em cada quadro do vídeo é feita com o classificador em "
        "cascata de Haar (Haar cascade) para faces frontais, considerando "
        "apenas o maior rosto encontrado, ou seja, o mais próximo da câmera. "
        "O rosto detectado é convertido para tons de cinza, redimensionado "
        "para 200 × 200 pixels e tem o histograma equalizado, para reduzir a "
        "influência da iluminação. O reconhecimento é feito com o algoritmo "
        "LBPH (Local Binary Patterns Histograms), que retorna a pessoa "
        "cadastrada mais parecida e um valor de distância (chamado de "
        "confiança no OpenCV): quanto menor esse valor, mais parecido é o "
        "rosto com o cadastrado.")
    par(f"Para evitar decisões baseadas em um único quadro, a presença só é "
        f"registrada quando a mesma pessoa é reconhecida em "
        f"{parametros['frames']} quadros seguidos com distância menor ou "
        f"igual ao limiar; quando um rosto não reconhecido aparece em "
        f"{parametros['frames']} quadros seguidos, o sistema informa que o "
        f"rosto não foi reconhecido. Se a pessoa já teve a presença "
        f"registrada no dia, uma nova passagem pela porta não gera registro "
        f"duplicado. A [[tab:parametros]] resume os parâmetros utilizados.")
    b.append(("tabela", "parametros", "Parâmetros do protótipo",
              ["Parâmetro", "Valor"], [
                  ["Detector de rosto",
                   "Haar cascade (haarcascade_frontalface_default)"],
                  ["Tamanho mínimo do rosto detectado", "80 × 80 pixels"],
                  ["Algoritmo de reconhecimento", "LBPH (OpenCV)"],
                  ["Pré-processamento",
                   "Tons de cinza, 200 × 200 px, equalização de histograma"],
                  ["Amostras por pessoa no cadastro", parametros["amostras"]],
                  ["Limiar de distância", num(parametros["limiar"], 1)],
                  ["Quadros seguidos para decidir", parametros["frames"]],
                  ["Tempo máximo por tentativa", f"{parametros['timeout']} s"],
              ], [6, 8]))

    titulo("Ambiente e equipamentos", 2)
    camera = config.get("camera") or "uma câmera de vídeo"
    par(f"A captura de vídeo foi feita com {camera}. O equipamento foi mantido "
        f"em posição fixa, simulando uma câmera instalada na porta de uma "
        f"sala de aula. O vídeo foi processado em um computador executando o "
        f"protótipo, que exibia na tela a imagem da câmera, a simulação da "
        f"porta e o resultado de cada tentativa.")

    titulo("Participantes", 2)
    papeis = {}
    for n in cadastrados:
        papeis.setdefault(pessoas[n], []).append(nome(n))
    desc_papeis = "; ".join(f"{p}: {', '.join(ns)}"
                            for p, ns in sorted(papeis.items()))
    texto = (f"Participaram dos testes {len(cadastrados)} voluntários, todos "
             f"cadastrados no sistema ({desc_papeis}). ")
    if config.get("anonimizar"):
        texto += ("Para preservar a identidade dos voluntários, eles são "
                  "identificados pelos códigos P1 a "
                  f"P{len(cadastrados)}. ")
    texto += (f"Cada participante foi cadastrado com {parametros['amostras']} "
              f"imagens do rosto capturadas pela própria câmera do protótipo.")
    if any(r["participante"] == VISITANTE for r in resultados):
        texto += (" No caso de teste de pessoa não cadastrada, foi utilizada "
                  "uma pessoa que não fazia parte do cadastro.")
    par(texto)

    titulo("Procedimento", 2)
    par(f"Para cada caso de teste e cada participante foram realizadas "
        f"{parametros['tentativas']} tentativas. Em cada tentativa, o operador "
        f"selecionava o caso e o participante na tela de testes, na qual as "
        f"decisões não são gravadas na lista de presença, e iniciava a "
        f"tentativa; o sistema então registrava automaticamente a decisão "
        f"(presença registrada ou não), a pessoa reconhecida, o tempo até a "
        f"decisão e os valores de distância do LBPH em cada quadro. Quando "
        f"não houve decisão em {parametros['timeout']} segundos, por exemplo "
        f"porque o rosto não foi detectado, a tentativa foi registrada como "
        f"\"sem decisão\" e considerada uma rejeição.")

    titulo("Métricas", 2)
    par("As tentativas foram divididas em genuínas, em que um participante "
        "cadastrado passa pela câmera com o próprio rosto e o esperado é o "
        "registro da sua presença, e de impostor, em que se apresenta uma "
        "foto, um vídeo ou uma pessoa não cadastrada e o esperado é que "
        "nenhuma presença seja registrada. Cada tentativa foi classificada "
        "como:")
    b.append(("itens", [
        ("Verdadeiro positivo (VP): ",
         "presença registrada para o próprio participante."),
        ("Falso negativo (FN): ",
         "participante presente sem presença registrada (rosto rejeitado ou "
         "não detectado)."),
        ("Falso positivo (FP): ",
         "presença registrada indevidamente, seja para um impostor, seja "
         "para outra pessoa que não o participante."),
        ("Verdadeiro negativo (VN): ",
         "tentativa de impostor sem presença registrada."),
    ]))
    par("A partir dessa classificação foram calculadas as métricas:")
    b.append(("itens", [
        ("Acurácia: ", "(VP + VN) / total de tentativas."),
        ("Taxa de registro correto (TAR): ", "VP / tentativas genuínas."),
        ("Taxa de falsa rejeição (FRR): ",
         "FN / tentativas genuínas, ou seja, a proporção de alunos presentes "
         "que ficariam com falta."),
        ("Taxa de erro de identificação: ",
         "presenças atribuídas a outra pessoa / tentativas genuínas."),
        ("Taxa de falsa aceitação (FAR): ",
         "impostores com presença registrada / tentativas de impostor."),
    ]))

    # ------------------------------------------------------------ casos
    titulo("Casos de teste")
    par(f"Foram definidos {len(CASOS)} casos de teste, apresentados na "
        f"[[tab:casos]]. Os casos CT01 a CT04 variam a distância entre o "
        f"rosto e a câmera, sendo o CT01 a condição de referência; os casos "
        f"CT05 a CT07 variam o ângulo do rosto; e os casos CT08 a CT10 "
        f"simulam tentativas de fraude.")
    b.append(("tabela", "casos", "Casos de teste",
              ["ID", "Grupo", "Descrição", "Procedimento", "Esperado"],
              [[c["id"], c["grupo"], c["titulo"], c["procedimento"],
                NOMES_DECISAO[c["esperado"]]] for c in CASOS],
              [1.2, 1.7, 3.3, 6.3, 2.2]))

    # ------------------------------------------------------------ resultados
    titulo("Resultados")
    if not resultados:
        par("Nenhuma tentativa foi registrada até o momento.")
        return b

    datas = sorted(r["data"][:10] for r in resultados)
    periodo = (_data_br(datas[0]) if datas[0] == datas[-1]
               else f"{_data_br(datas[0])} a {_data_br(datas[-1])}")

    titulo("Visão geral", 2)
    par(f"Os testes foram realizados em {periodo}, totalizando "
        f"{geral['total']} tentativas, das quais {geral['genuinas']} genuínas "
        f"e {geral['impostoras']} de impostor. A [[tab:geral]] apresenta as "
        f"métricas gerais. A acurácia geral foi de {pct(geral['acuracia'])}"
        + (f", com taxa de registro correto de {pct(geral['tar'])} e taxa de "
           f"falsa rejeição de {pct(geral['frr'])} nas tentativas genuínas"
           if geral["genuinas"] else "")
        + (f", e taxa de falsa aceitação de {pct(geral['far'])} nas "
           f"tentativas de impostor" if geral["impostoras"] else "")
        + ".")
    b.append(("tabela", "geral", "Métricas gerais", ["Métrica", "Valor"], [
        ["Total de tentativas", geral["total"]],
        ["Tentativas genuínas / de impostor",
         f"{geral['genuinas']} / {geral['impostoras']}"],
        ["VP / FN / FP / VN",
         f"{geral['VP']} / {geral['FN']} / {geral['FP']} / {geral['VN']}"],
        ["Acurácia", pct(geral["acuracia"])],
        ["Taxa de registro correto (TAR)", pct(geral["tar"])],
        ["Taxa de falsa rejeição (FRR)", pct(geral["frr"])],
        ["Taxa de erro de identificação", pct(geral["erro_id"])],
        ["Taxa de falsa aceitação (FAR)", pct(geral["far"])],
        ["Tentativas sem decisão", geral["sem_deteccao"]],
        ["Tempo médio até a decisão", f"{num(geral['tempo_medio'])} s"],
        ["Distância LBPH média", num(geral["confianca_media"], 1)],
    ], [7, 4]))

    titulo("Resultados por caso de teste", 2)
    par("A [[tab:por-caso]] detalha os resultados de cada caso de teste e a "
        "[[fig:acuracia]] compara a acurácia obtida em cada um deles.")
    b.append(("tabela", "por-caso", "Resultados por caso de teste",
              ["Caso", "N", "VP", "FN", "FP", "VN", "Acurácia",
               "Tempo médio (s)", "Dist. média"],
              [[c["id"], m["total"], m["VP"], m["FN"], m["FP"], m["VN"],
                pct(m["acuracia"]), num(m["tempo_medio"]),
                num(m["confianca_media"], 1)]
               for c in casos_feitos for m in [por_caso[c["id"]]]], None))
    b.append(("figura", "acuracia", "Acurácia por caso de teste",
              _grafico_barras(
                  [c["id"] for c in casos_feitos],
                  [("Acurácia (%)",
                    [por_caso[c["id"]]["acuracia"] for c in casos_feitos],
                    "#2a6fdb")],
                  "Acurácia (%)")))

    if genuinos:
        titulo("Variação de distância e ângulo", 2)
        ref = por_caso.get("CT01")
        melhor = max(genuinos, key=lambda c: por_caso[c["id"]]["tar"])
        pior = min(genuinos, key=lambda c: por_caso[c["id"]]["tar"])
        texto = ("A [[fig:genuinos]] mostra a taxa de registro correto e a "
                 "taxa de falsa rejeição nos casos genuínos. ")
        if ref:
            texto += (f"Na condição de referência (CT01, rosto frontal a "
                      f"50 cm), a taxa de registro correto foi de "
                      f"{pct(ref['tar'])}. ")
        texto += (f"O melhor desempenho foi observado no caso "
                  f"{melhor['id']} ({melhor['titulo'].lower()}), com "
                  f"{pct(por_caso[melhor['id']]['tar'])} de registro "
                  f"correto, e o pior no caso {pior['id']} "
                  f"({pior['titulo'].lower()}), com "
                  f"{pct(por_caso[pior['id']]['tar'])}. Na prática, cada "
                  f"falsa rejeição corresponde a um aluno presente que "
                  f"ficaria com falta e precisaria ter a presença lançada "
                  f"manualmente pelo professor.")
        par(texto)

        sem = [(c, por_caso[c["id"]]["sem_deteccao"]) for c in genuinos
               if por_caso[c["id"]]["sem_deteccao"]]
        if sem:
            lista = ", ".join(f"{c['id']} ({n})" for c, n in sem)
            par(f"Houve tentativas sem decisão nos casos {lista}. Nessas "
                f"tentativas o rosto não chegou a ser detectado de forma "
                f"contínua, o que é esperado em rostos muito distantes "
                f"(menores que o tamanho mínimo de 80 × 80 pixels do detector) "
                f"ou muito girados, já que o classificador de Haar utilizado "
                f"foi treinado para faces frontais.")
        erros_id = sum(por_caso[c["id"]]["FP"] for c in genuinos)
        if erros_id:
            par(f"Em {erros_id} tentativa(s) genuína(s) a presença foi "
                f"registrada para outra pessoa cadastrada. No contexto de uma "
                f"chamada, esse é o erro mais grave, pois tem efeito duplo: o "
                f"aluno presente fica com falta e outro aluno recebe uma "
                f"presença indevida.")
        else:
            par("Em nenhuma tentativa genuína a presença foi registrada para "
                "outra pessoa cadastrada.")
        b.append(("figura", "genuinos",
                  "Registro correto e falsa rejeição nos casos genuínos",
                  _grafico_barras(
                      [c["id"] for c in genuinos],
                      [("TAR (%)", [por_caso[c["id"]]["tar"] for c in genuinos],
                        "#2a9d5c"),
                       ("FRR (%)", [por_caso[c["id"]]["frr"] for c in genuinos],
                        "#d9534f")],
                      "Taxa (%)")))

    if fraudes:
        titulo("Tentativas de fraude", 2)
        partes = []
        for c in fraudes:
            m = por_caso[c["id"]]
            partes.append(f"no caso {c['id']} ({c['titulo'].lower()}), "
                          f"{m['FP']} de {m['total']} tentativas tiveram a "
                          f"presença registrada indevidamente (FAR de "
                          f"{pct(m['far'])})")
        par("Nas tentativas de impostor, " + "; ".join(partes) + ".")
        spoof = [c for c in fraudes if c["id"] in ("CT08", "CT09")]
        aceitos = sum(por_caso[c["id"]]["FP"] for c in spoof)
        if spoof and aceitos:
            par("A aceitação de fotos e vídeos exibidos na tela de um celular "
                "significa que, na prática, um aluno poderia registrar a "
                "presença de um colega ausente apenas mostrando uma imagem "
                "dele à câmera. Isso ocorre porque o protótipo não possui "
                "detecção de vivacidade (liveness detection): o LBPH compara "
                "apenas a textura da imagem do rosto em tons de cinza e não "
                "diferencia um rosto real de uma reprodução, uma limitação "
                "conhecida de sistemas baseados apenas em imagens 2D.")
        elif spoof:
            par("Nenhuma foto ou vídeo exibido no celular teve a presença "
                "registrada. Vale observar, porém, que o protótipo não possui "
                "uma etapa específica de detecção de vivacidade (liveness "
                "detection); a rejeição se deve provavelmente às diferenças "
                "que a tela introduz na imagem (reflexos, brilho, padrão "
                "moiré e perda de nitidez), que aumentam a distância LBPH. "
                "Com uma reprodução de melhor qualidade o resultado poderia "
                "ser diferente.")

    titulo("Resultados por participante", 2)
    par("A [[tab:participantes]] apresenta os resultados agrupados por "
        "participante.")
    b.append(("tabela", "participantes", "Resultados por participante",
              ["Participante", "N", "VP", "FN", "FP", "VN", "Acurácia",
               "Dist. média"],
              [[nome(p), m["total"], m["VP"], m["FN"], m["FP"], m["VN"],
                pct(m["acuracia"]), num(m["confianca_media"], 1)]
               for p, m in sorted(por_participante.items(),
                                  key=lambda kv: nome(kv[0]))], None))

    # ------------------------------------------------------------ conclusão
    titulo("Considerações sobre os resultados")
    texto = (f"Os testes mostraram que o protótipo atingiu acurácia geral de "
             f"{pct(geral['acuracia'])} nas {geral['total']} tentativas "
             f"realizadas. ")
    if genuinos:
        texto += (f"Nas tentativas genuínas, o melhor desempenho ocorreu no "
                  f"caso {melhor['id']} ({melhor['titulo'].lower()}) e o "
                  f"pior no caso {pior['id']} ({pior['titulo'].lower()}), o "
                  f"que deve ser levado em conta no posicionamento da câmera "
                  f"na porta da sala e na orientação dos alunos ao entrar. ")
    if fraudes:
        texto += (f"Nas tentativas de fraude, a taxa de falsa aceitação foi "
                  f"de {pct(geral['far'])}. ")
    texto += (
        "Como melhorias futuras, sugerem-se o uso de detectores e "
        "reconhecedores baseados em redes neurais, mais tolerantes a "
        "variações de ângulo e distância; a inclusão de detecção de "
        "vivacidade, para impedir o registro de presença com fotos e vídeos; "
        "a conferência da lista de presença pelo professor ao final da aula; "
        "e o ajuste do limiar de distância a partir de um número maior de "
        "participantes.")
    par(texto)

    # ------------------------------------------------------------ apêndice
    b.append(("apendice",))
    titulo("Apêndice A – Registro das tentativas")
    b.append(("tabela", "tentativas", "Registro de todas as tentativas",
              ["Caso", "Participante", "Nº", "Decisão", "Reconhecido como",
               "Classif.", "Tempo (s)", "Dist. média"],
              [[r["caso"], nome(r["participante"]), r["tentativa"],
                NOMES_DECISAO[r["decisao"]], nome(r["nome_reconhecido"]),
                r["classificacao"], num(r["tempo"]),
                num(r["confianca_media"], 1)]
               for r in sorted(resultados, key=lambda r: (
                   r["caso"], nome(r["participante"]), r["tentativa"]))],
              None))
    return b


# ---------------------------------------------------------------- .docx (ABNT)

def gerar_docx(config, parametros, resultados, pessoas, caminho=DOCX_PATH):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Cm, Pt, RGBColor

    blocos = montar_conteudo(config, parametros, resultados, pessoas)
    ano = time.strftime("%Y")
    doc = Document()

    # ABNT: A4, margens 3/2 cm, Arial 12, espaçamento 1,5
    for secao in doc.sections:
        secao.page_height, secao.page_width = Cm(29.7), Cm(21)
        secao.top_margin = secao.left_margin = Cm(3)
        secao.bottom_margin = secao.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    for nivel in (1, 2):
        estilo = doc.styles[f"Heading {nivel}"]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(12)
        estilo.font.bold = True
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        estilo.paragraph_format.space_before = Pt(18)
        estilo.paragraph_format.space_after = Pt(12)

    # numera tabelas e figuras na ordem em que aparecem
    numeros, cont = {}, {"tabela": 0, "figura": 0}
    for bl in blocos:
        if bl[0] in cont:
            cont[bl[0]] += 1
            numeros[("tab" if bl[0] == "tabela" else "fig", bl[1])] = cont[bl[0]]

    def resolver(texto):
        return REF.sub(lambda m: ("Tabela " if m[1] == "tab" else "Figura ")
                       + str(numeros[(m[1], m[2])]), texto)

    def centralizado(texto, negrito=False, tamanho=12):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto)
        r.bold = negrito
        r.font.size = Pt(tamanho)

    def legenda(texto):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(texto)
        r.font.size = Pt(10)
        r.bold = True

    def fonte():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        p.add_run(f"Fonte: elaborado pelo autor ({ano}).").font.size = Pt(10)

    def celula(c, texto, negrito=False):
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(texto))
        r.font.size = Pt(10)
        r.bold = negrito

    # cabeçalho
    if config.get("instituicao"):
        centralizado(config["instituicao"].upper(), negrito=True)
    if config.get("curso"):
        centralizado(config["curso"].upper())
    centralizado("")
    if config.get("titulo"):
        centralizado(config["titulo"].upper(), negrito=True, tamanho=14)
    centralizado("Relatório de testes do protótipo",
                 negrito=not config.get("titulo"))
    if config.get("autor"):
        centralizado(config["autor"])
    centralizado(", ".join(x for x in (config.get("cidade"), ano) if x))

    n1 = n2 = 0
    apendice = False
    for bl in blocos:
        tipo = bl[0]
        if tipo == "apendice":
            apendice = True
        elif tipo == "titulo":
            _, nivel, texto = bl
            if apendice:
                doc.add_heading(texto.upper(), 1)
            elif nivel == 1:
                n1, n2 = n1 + 1, 0
                doc.add_heading(f"{n1} {texto.upper()}", 1)
            else:
                n2 += 1
                doc.add_heading(f"{n1}.{n2} {texto}", 2)
        elif tipo == "par":
            p = doc.add_paragraph(resolver(bl[1]))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.25)
        elif tipo == "itens":
            for rotulo, texto in bl[1]:
                p = doc.add_paragraph(style="List Bullet")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.add_run(rotulo).bold = True
                p.add_run(texto)
        elif tipo == "tabela":
            _, id_, titulo, cabecalho, linhas, larguras = bl
            legenda(f"Tabela {numeros[('tab', id_)]} – {titulo}")
            t = doc.add_table(rows=1, cols=len(cabecalho))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, texto in enumerate(cabecalho):
                celula(t.rows[0].cells[i], texto, negrito=True)
            for linha in linhas:
                celulas = t.add_row().cells
                for i, texto in enumerate(linha):
                    celula(celulas[i], texto)
            if larguras:
                for linha in t.rows:
                    for i, w in enumerate(larguras):
                        linha.cells[i].width = Cm(w)
            fonte()
        elif tipo == "figura":
            _, id_, titulo, png = bl
            legenda(f"Figura {numeros[('fig', id_)]} – {titulo}")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(io.BytesIO(png), width=Cm(15))
            fonte()

    doc.save(caminho)
    return caminho


# ---------------------------------------------------------------- LaTeX (SBC)

_ESCAPES = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$",
            "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}",
            "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
            "°": r"$^\circ$", "×": r"$\times$", "–": "--"}


def _tex(texto):
    texto = "".join(_ESCAPES.get(ch, ch) for ch in str(texto))
    texto = re.sub(r'"([^"]*)"', r"``\1''", texto)  # aspas tipográficas
    return REF.sub(lambda m: ("Tabela~" if m[1] == "tab" else "Figura~")
                   + rf"\ref{{{m[1]}:{m[2]}}}", texto)


def _tabela_latex(id_, titulo, cabecalho, linhas, larguras):
    if larguras:
        # cabe na largura do texto do template SBC (~16 cm)
        fator = min(1.0, 14.5 / sum(larguras))
        colunas = "|".join(rf"p{{{w * fator:.1f}cm}}" for w in larguras)
    else:
        colunas = "|".join(["l"] + ["c"] * (len(cabecalho) - 1))
    linhas_tex = [" & ".join(rf"\textbf{{{_tex(c)}}}" for c in cabecalho)]
    linhas_tex += [" & ".join(_tex(c) for c in linha) for linha in linhas]
    corpo = " \\\\ \\hline\n".join(linhas_tex)
    return (f"\\begin{{table}}[ht]\n\\centering\n"
            f"\\caption{{{_tex(titulo)}}}\n\\label{{tab:{id_}}}\n\\small\n"
            f"\\begin{{tabular}}{{|{colunas}|}}\n\\hline\n"
            f"{corpo} \\\\ \\hline\n\\end{{tabular}}\n\\end{{table}}\n")


def gerar_latex(config, parametros, resultados, pessoas, caminho=LATEX_PATH):
    """Gera um .zip com secao_testes.tex + figuras, para o template da SBC.

    Os títulos viram \\subsection/\\subsubsection dentro de uma
    \\section{Avaliação Experimental}; o apêndice com todas as tentativas
    fica de fora (artigo tem limite de páginas; os dados estão no CSV).
    """
    blocos = montar_conteudo(config, parametros, resultados, pessoas)
    partes = ["% Seção gerada automaticamente pelo protótipo (relatorio.py).",
              "% No artigo: \\input{secao_testes}",
              "",
              "\\section{Avaliação Experimental}\\label{sec:avaliacao}",
              ""]
    figuras = {}
    for bl in blocos:
        tipo = bl[0]
        if tipo == "apendice":
            break
        if tipo == "titulo":
            comando = "subsection" if bl[1] == 1 else "subsubsection"
            partes.append(f"\\{comando}{{{_tex(bl[2])}}}\n")
        elif tipo == "par":
            partes.append(_tex(bl[1]) + "\n")
        elif tipo == "itens":
            itens = "\n".join(rf"  \item \textbf{{{_tex(r.strip())}}} {_tex(t)}"
                              for r, t in bl[1])
            partes.append(f"\\begin{{itemize}}\n{itens}\n\\end{{itemize}}\n")
        elif tipo == "tabela":
            partes.append(_tabela_latex(*bl[1:]))
        elif tipo == "figura":
            _, id_, titulo, png = bl
            arquivo = f"figuras/{id_}.png"
            figuras[arquivo] = png
            partes.append(
                f"\\begin{{figure}}[ht]\n\\centering\n"
                f"\\includegraphics[width=.9\\textwidth]{{{arquivo}}}\n"
                f"\\caption{{{_tex(titulo)}}}\n\\label{{fig:{id_}}}\n"
                f"\\end{{figure}}\n")

    leiame = (
        "Como usar no artigo (template SBC)\n"
        "==================================\n\n"
        "1. Copie secao_testes.tex e a pasta figuras/ para a pasta do seu .tex.\n"
        "2. No ponto do artigo onde a seção deve entrar, coloque:\n"
        "       \\input{secao_testes}\n"
        "3. O arquivo está em UTF-8. Confira se o preâmbulo do artigo tem\n"
        "       \\usepackage[utf8]{inputenc}\n"
        "   (versões antigas do template usam latin1: troque para utf8).\n"
        "4. O template da SBC já carrega o pacote graphicx, usado nas figuras.\n"
        "5. As referências usam os rótulos tab:... e fig:...; para citar a\n"
        "   seção em outra parte do artigo use \\ref{sec:avaliacao}.\n\n"
        "A tabela com todas as tentativas não foi incluída (o artigo tem\n"
        "limite de páginas); ela está no CSV exportado pela página de testes\n"
        "e no apêndice do relatório .docx.\n")

    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("secao_testes.tex", "\n".join(partes))
        z.writestr("LEIAME.txt", leiame)
        for arquivo, png in figuras.items():
            z.writestr(arquivo, png)
    return caminho


if __name__ == "__main__":
    import sys
    from app import (carregar_pessoas, carregar_relatorio_config,
                     parametros_do_prototipo)

    gerar = gerar_latex if "latex" in sys.argv[1:] else gerar_docx
    arquivo = gerar(carregar_relatorio_config(), parametros_do_prototipo(),
                    carregar_resultados(), carregar_pessoas())
    print(f"Relatório gerado: {arquivo}")
